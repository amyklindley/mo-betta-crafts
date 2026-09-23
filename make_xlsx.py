"""Builds MoBetta-Recipes.xlsx: one tab per profession from recipes.json."""
import json, datetime, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open("recipes.json", encoding="utf-8"))
rows = d if isinstance(d, list) else d.get("recipes", d)
if isinstance(rows, dict): rows = list(rows.values())
rows = [json.loads(json.dumps(r).replace("\ufffd", "-").replace("\u2013", "-")) for r in rows]
by_skill = {}
for r in rows: by_skill.setdefault(r["skill"], []).append(r)

ARIAL = "Arial"
hdr_font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="3B3F4A")
body_font = Font(name=ARIAL, size=10)
link_font = Font(name=ARIAL, size=10, color="0563C1", underline="single")
thin = Side(style="thin", color="D9D9D9")
border = Border(bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
top = Alignment(vertical="top")
center = Alignment(horizontal="center", vertical="top")

def ing_text(r):
    parts = []
    for i in r["ingredients"]:
        s = f"{i['qty']} x {i['name']}" if i.get("qty", 1) != 1 else i["name"]
        if i.get("tool"): s += " (tool)"
        parts.append(s)
    return "\n".join(parts)

def res_text(r):
    return "\n".join(f"{x['qty']} x {x['name']}" if x.get("qty", 1) != 1 else x["name"] for x in r["results"]) or r["name"]

headers = ["Recipe", "Makes", "Trivial", "Station", "Section", "Ingredients", "Result(s)", "Notes", "Salvage", "Wiki"]
widths  = [34, 7, 8, 16, 30, 44, 30, 30, 8, 12]

wb = Workbook()
idx = wb.active; idx.title = "Index"
skills = sorted(by_skill, key=lambda s: -len(by_skill[s]))
safe = lambda s: re.sub(r"[\[\]\*\?/\:]", "", s)[:31]

for skill in skills:
    ws = wb.create_sheet(safe(skill))
    ws.append(headers)
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=c); cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    recs = sorted(by_skill[skill], key=lambda r: (r["trivial"] if isinstance(r["trivial"], (int, float)) else 10**6, r["section"], r["name"]))
    for r in recs:
        ws.append([r["name"], r["makes"], r["trivial"] if r["trivial"] not in ("", None) else None, r["station"], r["section"],
                   ing_text(r), res_text(r), r["notes"], "yes" if r["salvage"] else "", "wiki"])
        rr = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=rr, column=c); cell.font = body_font; cell.border = border
            cell.alignment = wrap if c in (1, 5, 6, 7, 8) else (center if c in (2, 3, 9) else top)
        link = ws.cell(row=rr, column=10); link.hyperlink = r["url"]; link.font = link_font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

idx["A1"] = "Mo Betta Crafts: Monsters & Memories recipes"; idx["A1"].font = Font(name=ARIAL, bold=True, size=14)
idx["A2"] = f"Source: Monsters and Memories Wiki (monstersandmemories.miraheze.org), Skill pages, scraped {datetime.date.today():%Y-%m-%d}. One tab per profession; sorted by trivial, then section, then name."
idx["A2"].font = Font(name=ARIAL, size=9, italic=True, color="666666")
idx["A4"], idx["B4"], idx["C4"] = "Profession", "Recipes", "Tab"
for c in (1, 2, 3):
    cell = idx.cell(row=4, column=c); cell.font = hdr_font; cell.fill = hdr_fill
for i, skill in enumerate(skills, start=5):
    idx.cell(row=i, column=1, value=skill).font = body_font
    f = idx.cell(row=i, column=2, value=f"=COUNTA('{safe(skill)}'!A:A)-1"); f.font = body_font; f.alignment = Alignment(horizontal="right")
    l = idx.cell(row=i, column=3, value="open tab"); l.hyperlink = f"#'{safe(skill)}'!A1"; l.font = link_font
tot = 5 + len(skills)
idx.cell(row=tot, column=1, value="Total").font = Font(name=ARIAL, bold=True, size=10)
t = idx.cell(row=tot, column=2, value=f"=SUM(B5:B{tot-1})"); t.font = Font(name=ARIAL, bold=True, size=10); t.alignment = Alignment(horizontal="right")
idx.cell(row=tot + 2, column=1, value="Columns: Trivial = skill level at which the combine stops giving skill-ups (blank where the wiki did not list it). Station = where you combine it. Salvage = the recipe breaks an item down rather than making one. Ingredients marked (tool) are not consumed.").font = Font(name=ARIAL, size=9, color="666666")
idx.column_dimensions["A"].width = 22; idx.column_dimensions["B"].width = 10; idx.column_dimensions["C"].width = 12
idx.freeze_panes = "A5"
wb.save("MoBetta-Recipes.xlsx"); print("saved | tabs:", [safe(s) for s in skills], "| rows:", len(rows))
